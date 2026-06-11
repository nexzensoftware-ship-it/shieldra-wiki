from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.shared import Inches, Pt


OUT = "Microsoft_Comparative_Analysis_APA7.docx"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "24")
    rpr.append(sz)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    fld.append(run)
    p._p.append(fld)


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)

    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.line_spacing = 2.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Inches(0)


def add_para(doc, text, style=None, indent=True, align=None):
    p = doc.add_paragraph(style=style)
    if not indent:
        p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    p.add_run(text)
    return p


def add_heading(doc, text, level=1, centered=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 2.0
    p.paragraph_format.first_line_indent = Inches(0)
    r = p.add_run("Comparative Analysis of Microsoft Corporation's Business Performance")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    for text in [
        "[Student Name]",
        "[University Name]",
        "[Course Number and Name]",
        "[Instructor Name]",
        "May 4, 2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 2.0
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def add_comparison_table(doc):
    add_para(doc, "Table 1", indent=False)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Inches(0)
    p.paragraph_format.line_spacing = 2.0
    r = p.add_run("Comparative performance summary")
    r.italic = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    p.paragraph_format.keep_with_next = True

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.35), Inches(2.0), Inches(2.0), Inches(2.15)]
    headers = ["Category", "Microsoft evidence", "Benchmark or competitor", "Strategic implication"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, h, bold=True)
        set_cell_shading(cell, "EDEDED")

    rows = [
        [
            "Finance",
            "FY2025 revenue of $281.7 billion, operating income of $128.5 billion, net income of $101.8 billion, and $168.9 billion in Microsoft Cloud revenue.",
            "AWS remained the largest cloud infrastructure provider at 28% Q4 2025 market share; Microsoft held 21%.",
            "Microsoft has outstanding profitability and cash generation, but continued cloud share gains require careful AI capital allocation.",
        ],
        [
            "Innovation and AI",
            "Azure revenue surpassed $75 billion and grew 34%; R&D was $32.5 billion, or 12% of revenue.",
            "Google Cloud reached $13.9 billion in 2025 operating income; AWS generated $45.6 billion in AWS operating income.",
            "Microsoft's AI stack is broad and commercially advanced, but rivals are investing quickly and customers need proof of measurable AI value.",
        ],
        [
            "Sustainability and operations",
            "Microsoft contracted nearly 22 million metric tons of carbon removals in FY2024 and expanded water reuse projects.",
            "Google reported data-center energy emissions reductions and Amazon reported continued renewable energy procurement while AI demand rose across the sector.",
            "Microsoft's transparency is a strength, yet Scope 3 emissions and water intensity remain material improvement areas.",
        ],
    ]

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].width = widths[i]
            set_cell_text(cells[i], text)

    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        for cell in row.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_mar = OxmlElement("w:tcMar")
            for side in ["top", "left", "bottom", "right"]:
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), "90")
                node.set(qn("w:type"), "dxa")
                tc_mar.append(node)
            tc_pr.append(tc_mar)


def add_references(doc):
    doc.add_page_break()
    add_heading(doc, "References", level=1, centered=True)
    refs = [
        "Alphabet Inc. (2026). Annual report on Form 10-K for the fiscal year ended December 31, 2025. https://s206.q4cdn.com/479360582/files/doc_financials/2025/q4/GOOG-10-K-2025.pdf",
        "Amazon.com, Inc. (2026). 2025 annual report. https://s2.q4cdn.com/299287126/files/doc_financials/2026/ar/Amazon-2025-Annual-Report.pdf",
        "Casey, K. (2026, February 6). GenAI drives $119B cloud revenue in Q4. TechTarget. https://www.techtarget.com/searchcloudcomputing/news/366638805/GenAI-drives-119B-cloud-revenue-in-Q4",
        "Google. (2025). 2025 environmental report. https://www.sustainability.google/reports/google-2025-environmental-report/",
        "Microsoft. (2025a). Annual report 2025. https://www.microsoft.com/investor/reports/ar25/",
        "Microsoft. (2025b, May 29). Our 2025 Environmental Sustainability Report. Microsoft On the Issues. https://blogs.microsoft.com/on-the-issues/2025/05/29/environmental-sustainability-report/",
        "Microsoft. (2025c). 2025 Environmental Sustainability Report. https://www.microsoft.com/en-us/corporate-responsibility/sustainability/report/",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 2.0
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(ref)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)


def build():
    doc = Document()
    style_document(doc)
    add_title_page(doc)
    doc.add_page_break()

    add_heading(doc, "Comparative Analysis of Microsoft Corporation's Business Performance", centered=True)
    add_para(
        doc,
        "Microsoft Corporation is a well-known publicly traded technology company whose mission is to empower every person and every organization on the planet to achieve more. That mission gives Microsoft a useful lens for business performance analysis because the company's strategy depends on making enterprise technology, productivity software, cloud infrastructure, artificial intelligence, gaming, and security valuable at global scale. This paper evaluates Microsoft in three business categories: financial performance, innovation and artificial intelligence strategy, and sustainability and operations. The analysis identifies where the company excels, where improvement is needed, how Microsoft compares with competitors or industry benchmarks, and what recommendations would strengthen performance while remaining aligned with its mission and strategic goals.",
    )
    doc.add_page_break()
    add_comparison_table(doc)
    doc.add_page_break()

    add_heading(doc, "Financial Performance", level=1)
    add_heading(doc, "Where Microsoft Excels", level=2)
    add_para(
        doc,
        "Microsoft's financial performance is the company's clearest area of strength. In fiscal year 2025, Microsoft reported revenue of $281.7 billion, a 15% increase from the prior year, and operating income of $128.5 billion, a 17% increase (Microsoft, 2025a). Net income reached $101.8 billion, compared with $88.1 billion in fiscal year 2024, showing that growth translated into substantial bottom-line performance rather than only top-line expansion (Microsoft, 2025a). The company's profitability is supported by a diversified portfolio that includes Microsoft 365, Azure, LinkedIn, Dynamics, Windows, gaming, search advertising, and security products. This diversification reduces dependence on any single product cycle and gives Microsoft multiple ways to deepen customer relationships.",
    )
    add_para(
        doc,
        "The company's cloud business is especially important. Microsoft Cloud revenue was $168.9 billion in fiscal year 2025, up from $137.7 billion in 2024 and $111.6 billion in 2023 (Microsoft, 2025a). Azure surpassed $75 billion in annual revenue for the first time and grew 34%, reflecting broad demand for cloud and artificial intelligence workloads (Microsoft, 2025a). Microsoft also ended fiscal year 2025 with $94.6 billion in cash, cash equivalents, and short-term investments, giving the company significant flexibility to fund capital expenditures, research and development, acquisitions, dividends, and share repurchases (Microsoft, 2025a).",
    )
    add_heading(doc, "Where Improvement Is Needed", level=2)
    add_para(
        doc,
        "The main financial concern is not weak demand; it is the cost of meeting that demand. Microsoft acknowledged that investments in cloud and AI infrastructure will continue to increase operating costs and may reduce operating margins (Microsoft, 2025a). The company's Microsoft Cloud gross margin decreased to 69% because of the costs of scaling AI infrastructure, even though Azure efficiency gains partly offset the pressure (Microsoft, 2025a). This issue matters because investors and customers are watching whether generative AI creates durable economic value or only expensive usage growth. If AI capacity is built faster than monetized demand, the company could face lower returns on invested capital.",
    )
    add_heading(doc, "Comparative Assessment", level=2)
    add_para(
        doc,
        "Compared with the broader cloud industry, Microsoft is both a leader and a challenger. Synergy Research Group data reported by TechTarget estimated that Q4 2025 cloud infrastructure revenue reached $119.1 billion and that the three largest providers accounted for approximately two-thirds of enterprise spending; Amazon Web Services led with 28% share, Microsoft held 21%, and Google held 15% (Casey, 2026). This comparison shows that Microsoft is a global cloud leader but still trails AWS in market share. Amazon's 2025 annual report showed AWS operating income of $45.6 billion, while Alphabet reported Google Cloud operating income of $13.9 billion for 2025 (Alphabet Inc., 2026; Amazon.com, Inc., 2026). Microsoft's Intelligent Cloud operating income of $44.6 billion is therefore competitive with AWS, although the segment is not identical to AWS because Microsoft's segment includes additional server products and enterprise services (Microsoft, 2025a).",
    )
    add_heading(doc, "Recommendations", level=2)
    add_para(
        doc,
        "Microsoft should strengthen financial performance by creating more transparent internal and external measures of AI investment returns. Management should track AI infrastructure spending by workload type, customer segment, utilization rate, gross margin, and renewal impact. The company should also continue optimizing its mix of purchased GPUs, internally designed infrastructure, and partner capacity to avoid overdependence on scarce components. These recommendations support Microsoft's mission because financially sustainable AI capacity allows more organizations to access useful technology at scale without forcing the company to sacrifice long-term profitability.",
    )

    add_heading(doc, "Innovation and Artificial Intelligence Strategy", level=1)
    add_heading(doc, "Where Microsoft Excels", level=2)
    add_para(
        doc,
        "Microsoft excels in innovation because it has embedded AI across a full commercial technology stack rather than treating it as a stand-alone product. The company connects Azure infrastructure, Azure AI services, Microsoft 365 Copilot, GitHub Copilot, Dynamics 365, security tools, Windows, and developer platforms. This integrated approach helps Microsoft convert research and engineering investments into products that customers already use. Fiscal year 2025 research and development expense was $32.5 billion, or 12% of revenue, and increased 10% because of investments in cloud and AI engineering and gaming (Microsoft, 2025a).",
    )
    add_para(
        doc,
        "The company's innovation advantage is also visible in commercialization. Azure's 34% growth and Microsoft Cloud revenue of $168.9 billion show that Microsoft is not merely experimenting with AI; it is selling AI-enabled infrastructure, data platforms, productivity software, and developer tools at global scale (Microsoft, 2025a). The company benefits from enterprise trust, identity management, security capabilities, and established procurement relationships. Those factors make it easier for customers to adopt Copilot or Azure AI in regulated and complex environments than it would be for a newer AI company without enterprise distribution.",
    )
    add_heading(doc, "Where Improvement Is Needed", level=2)
    add_para(
        doc,
        "The improvement area is customer value realization. Microsoft has strong AI products, but enterprise buyers increasingly expect evidence that AI improves productivity, reduces costs, improves quality, or accelerates revenue. Broad adoption can slow if customers view AI tools as costly add-ons rather than measurable business investments. Microsoft also faces execution risk from rapid product expansion. Too many overlapping Copilot, Azure AI, and partner offerings could create confusion for customers deciding which tools to deploy, govern, secure, and measure.",
    )
    add_heading(doc, "Comparative Assessment", level=2)
    add_para(
        doc,
        "Microsoft's AI strategy compares favorably with AWS and Google because it combines cloud infrastructure with productivity software and developer workflows. AWS remains the largest cloud infrastructure provider and reported AWS operating income of $45.6 billion in 2025, demonstrating the profitability of its infrastructure leadership (Amazon.com, Inc., 2026). Google Cloud, by contrast, is growing rapidly and increased operating income to $13.9 billion in 2025, supported by Google infrastructure, AI models, and enterprise cloud demand (Alphabet Inc., 2026). Microsoft's advantage is that AI can be monetized through both Azure consumption and software subscriptions, while its gap is that AWS still has the larger cloud infrastructure share and Google has deep in-house AI research capabilities.",
    )
    add_heading(doc, "Recommendations", level=2)
    add_para(
        doc,
        "Microsoft should simplify AI adoption by publishing clearer industry-specific value frameworks for Copilot and Azure AI. These frameworks should define expected use cases, data governance requirements, productivity metrics, implementation timelines, and cost ranges for industries such as health care, financial services, manufacturing, education, and government. The company should also expand customer success programs that help organizations redesign workflows instead of simply licensing AI tools. This recommendation aligns with Microsoft's mission because customers achieve more when innovation is translated into reliable business outcomes.",
    )

    add_heading(doc, "Sustainability and Operations", level=1)
    add_heading(doc, "Where Microsoft Excels", level=2)
    add_para(
        doc,
        "Microsoft excels in sustainability transparency and in using its scale to develop new environmental solutions. The company has committed to becoming carbon negative, water positive, and zero waste by 2030, and its reporting openly acknowledges progress as well as setbacks. Microsoft reported that it contracted nearly 22 million metric tons of carbon removals in fiscal year 2024, and it highlighted water reuse work in Quincy, Washington, that reduced potable water use for cooling in that region by 97% while providing 1.5 million cubic meters of water annually for community drinking needs (Microsoft, 2025b; Microsoft, 2025c). These examples show that Microsoft's sustainability strategy is linked to operations rather than only public relations.",
    )
    add_para(
        doc,
        "The company also uses partnerships to address difficult supply chain emissions. For example, Microsoft worked with Elemental Impact, Capgemini, and Bouygues to launch the Build Better Innovation Challenge, which targets low-carbon materials that could reduce Scope 3 emissions from construction and infrastructure (Microsoft, 2025c). This approach is strategically relevant because data-center growth depends on steel, concrete, servers, chips, energy, and cooling systems. Improvements in these areas can reduce environmental impact while also improving operational resilience.",
    )
    add_heading(doc, "Where Improvement Is Needed", level=2)
    add_para(
        doc,
        "The main sustainability challenge is that Microsoft's emissions have risen even as the company invests heavily in decarbonization. Microsoft reported that total Scope 1, 2, and 3 emissions increased 23.4% compared with its 2020 baseline, largely because of AI and cloud expansion, even though revenue grew 71% and energy use rose 168% over the same period (Microsoft, 2025b). Scope 3 emissions are especially difficult because they include data-center construction, hardware manufacturing, suppliers, logistics, and other value-chain activities. The company therefore needs faster progress on supplier decarbonization, lower-carbon construction materials, circular hardware, and location decisions that account for energy and water constraints.",
    )
    add_heading(doc, "Comparative Assessment", level=2)
    add_para(
        doc,
        "Microsoft is not alone in facing this tension. AI has increased infrastructure demand across the technology sector. Google reported that its 2025 environmental report focused heavily on AI, energy, and resilience, and stated that its data centers used 84% less overhead energy than the industry average in 2024 (Google, 2025). Amazon also reported continued renewable energy procurement and data-center sustainability innovation while AWS demand expanded (Amazon.com, Inc., 2026). Microsoft's advantage is its unusually direct disclosure of the conflict between AI growth and sustainability goals. Its gap is that transparency alone does not close the emissions curve; performance must eventually show absolute reductions consistent with the 2030 targets.",
    )
    add_heading(doc, "Recommendations", level=2)
    add_para(
        doc,
        "Microsoft should make sustainability a formal constraint in AI capacity planning. Data-center investment decisions should include carbon intensity, water stress, grid readiness, local community impact, and supplier emissions alongside cost, latency, and customer demand. The company should also require major suppliers to use renewable energy and low-carbon materials on a faster timeline, supported by long-term purchase agreements that make cleaner materials commercially viable. Finally, Microsoft should publish more granular progress metrics for AI-related emissions, water use, and hardware circularity. These actions align with the company's mission because technology that empowers every organization should not undermine the environmental conditions those organizations depend on.",
    )

    add_heading(doc, "Conclusion", level=1)
    add_para(
        doc,
        "Microsoft's overall performance is strong, but the company's future depends on disciplined execution. Financially, Microsoft combines rapid revenue growth, excellent profitability, substantial liquidity, and a highly scalable cloud business. In innovation, the company has built one of the most complete AI commercialization platforms in the industry by connecting infrastructure, productivity software, developer tools, and enterprise trust. In sustainability and operations, Microsoft has credible goals, strong transparency, and meaningful partnerships, but its emissions growth shows that the hardest work remains ahead.",
    )
    add_para(
        doc,
        "The analysis indicates that Microsoft should focus on three priorities: improve AI investment return measurement, simplify customer AI adoption around measurable business outcomes, and integrate sustainability constraints into data-center and supplier decisions. These recommendations are consistent with Microsoft's mission and strategic goals because they make growth more durable, useful, and responsible. Continuous improvement is especially important for a company of Microsoft's size because its decisions influence customers, suppliers, energy systems, software developers, and the broader technology industry. Microsoft is already a high-performing company, but its next competitive advantage will come from proving that AI growth, financial discipline, and environmental responsibility can advance together.",
    )

    add_references(doc)
    doc.save(OUT)


if __name__ == "__main__":
    build()
